#!/usr/bin/env python3
"""Build .docx, .pdf, and .txt from a resume markdown file.

Usage: .venv/bin/python tools/build_resume.py path/to/Resume.md [--no-pdf]

Format (see profile/writing_rules.md):
  # Name                       -> name
  contact line                 -> line right after the name
  ## Section                   -> section heading
  ### a | b | c | dates        -> entry line (last field right-aligned)
  - bullet                     -> bullet (supports **bold**)
  plain text                   -> paragraph
Requires: python-docx; LibreOffice (soffice) for PDF.
"""
import re
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HEAD_FONT = "Liberation Sans"   # LibreOffice substitutes if not installed
BODY_FONT = "Liberation Sans"
BODY_PT = 10.5
ACCENT = RGBColor(0x1A, 0x10, 0x40)
TEXT_WIDTH = Inches(7.3)

FORBIDDEN = ["—"]  # em dash


def set_font(run, name, size, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)
    if color is not None:
        run.font.color.rgb = color


def add_rich(par, text, size=BODY_PT):
    for i, part in enumerate(re.split(r"\*\*", text)):
        if part:
            set_font(par.add_run(part), BODY_FONT, size, bold=(i % 2 == 1))


def bottom_border(par):
    ppr = par._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), "6")
    b.set(qn("w:space"), "1")
    b.set(qn("w:color"), "1A1040")
    pbdr.append(b)
    ppr.append(pbdr)


def spacing(par, before=0, after=0):
    par.paragraph_format.space_before = Pt(before)
    par.paragraph_format.space_after = Pt(after)
    par.paragraph_format.line_spacing = 1.05


def build_docx(lines, out):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Inches(0.6))
    sec.top_margin = sec.bottom_margin = Inches(0.5)
    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = Pt(BODY_PT)

    expect_contact = False
    for raw in lines:
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("# "):
            p = doc.add_paragraph()
            spacing(p, 0, 2)
            set_font(p.add_run(line[2:].strip()), HEAD_FONT, 20, bold=True, color=ACCENT)
            expect_contact = True
            continue
        if expect_contact and not line.startswith("## "):
            p = doc.add_paragraph()
            if "|" in line:
                spacing(p, 0, 4)
                set_font(p.add_run(line.strip()), BODY_FONT, 9.5)
            else:  # headline / title line under the name
                spacing(p, 0, 1)
                set_font(p.add_run(line.strip()), BODY_FONT, 11.5, bold=True, color=ACCENT)
            continue
        expect_contact = False
        if line.startswith("## "):
            p = doc.add_paragraph()
            spacing(p, 8, 3)
            set_font(p.add_run(line[3:].strip().upper()), HEAD_FONT, 11.5, bold=True, color=ACCENT)
            bottom_border(p)
            p.paragraph_format.keep_with_next = True
            continue
        if line.startswith("### "):
            fields = [f.strip() for f in line[4:].split("|")]
            has_dates = len(fields) >= 2 and re.search(r"\d{4}|Present", fields[-1])
            left = fields[:-1] if has_dates else fields
            if not has_dates:
                p = doc.add_paragraph()
                spacing(p, 4, 1)
                p.paragraph_format.keep_with_next = True
                set_font(p.add_run(left[0]), BODY_FONT, BODY_PT + 0.5, bold=True)
                if len(left) > 1:
                    set_font(p.add_run("  |  " + "  |  ".join(left[1:])), BODY_FONT, BODY_PT)
                continue
            tbl = doc.add_table(rows=1, cols=2)
            tbl.autofit = False
            widths = (Inches(5.5), Inches(1.8))
            for col, cell, w in zip(tbl.columns, tbl.rows[0].cells, widths):
                col.width = w
                cell.width = w
            tblpr = tbl._tbl.tblPr
            layout = OxmlElement("w:tblLayout")
            layout.set(qn("w:type"), "fixed")
            tblpr.append(layout)
            for gc, w in zip(tbl._tbl.tblGrid.findall(qn("w:gridCol")), widths):
                gc.set(qn("w:w"), str(int(w.inches * 1440)))
            lp = tbl.rows[0].cells[0].paragraphs[0]
            rp = tbl.rows[0].cells[1].paragraphs[0]
            for par in (lp, rp):
                spacing(par, 4, 1)
                par.paragraph_format.keep_with_next = True
            rp.alignment = 2  # right
            set_font(lp.add_run(left[0]), BODY_FONT, BODY_PT + 0.5, bold=True)
            if len(left) > 1:
                set_font(lp.add_run("  |  " + "  |  ".join(left[1:])), BODY_FONT, BODY_PT)
            if has_dates:
                set_font(rp.add_run(fields[-1].replace(" ", "\u00a0")), BODY_FONT, BODY_PT)
            continue
        if line.lstrip().startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            spacing(p, 0, 1)
            p.paragraph_format.left_indent = Inches(0.22)
            add_rich(p, line.lstrip()[2:].strip())
            continue
        p = doc.add_paragraph()
        spacing(p, 0, 2)
        add_rich(p, line.strip())
    doc.save(out)


def build_txt(lines, out):
    res, contact = [], False
    for raw in lines:
        line = raw.rstrip()
        if line.startswith("# "):
            res.append(line[2:].strip()); contact = True; continue
        if contact and line.strip():
            res.append(line.strip()); contact = False; continue
        if line.startswith("## "):
            res += ["", line[3:].strip().upper()]; continue
        if line.startswith("### "):
            res += ["", " | ".join(f.strip() for f in line[4:].split("|"))]; continue
        if line.lstrip().startswith("- "):
            res.append("- " + line.lstrip()[2:].replace("**", "").strip()); continue
        if line.strip():
            res.append(line.replace("**", "").strip())
    Path(out).write_text("\n".join(res).strip() + "\n", encoding="utf-8")


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    src = Path(sys.argv[1])
    text = src.read_text(encoding="utf-8")
    for ch in FORBIDDEN:
        if ch in text:
            sys.exit(f"Refusing to build: forbidden character {ch!r} (em dash) found in {src}")
    # profile/banned_terms.txt: one term per line (case-insensitive), # for comments.
    # Setup fills it with names the candidate never wants on a resume.
    banned = Path(__file__).resolve().parent.parent / "profile" / "banned_terms.txt"
    if banned.exists():
        for term in banned.read_text(encoding="utf-8").splitlines():
            term = term.strip()
            if term and not term.startswith("#") and term.lower() in text.lower():
                sys.exit(f"Refusing to build: banned term {term!r} (profile/banned_terms.txt) found in {src}")
    lines = text.splitlines()
    docx_out, txt_out = src.with_suffix(".docx"), src.with_suffix(".txt")
    build_docx(lines, docx_out)
    build_txt(lines, txt_out)
    print(f"Built {docx_out}\nBuilt {txt_out}")
    if "--no-pdf" in sys.argv:
        return
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("LibreOffice not found: skipped PDF. Upload the .docx instead.")
        return
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir",
                    str(src.parent), str(docx_out)], check=True,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    print(f"Built {src.with_suffix('.pdf')}")


if __name__ == "__main__":
    main()
